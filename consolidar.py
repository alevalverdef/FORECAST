"""Une los tres Excel de la JPS en una sola base con un resultado por fecha.

- Histórico general.xlsx: 'dd de Mes aaaa - Día' | número. El tipo de sorteo se infiere:
  domingo = Lotería Nacional, cualquier otro día = Chances.
- Histórico loteria.xlsx / Histórico Chances.xlsx: celdas combinadas (2 filas por sorteo);
  la fecha puede venir como fecha de Excel o como texto 'dd Mon, aaaa'. En Chances el
  bloque D:F es una copia desplazada de A:C, así que se lee solo A:C.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pandas as pd

MESES = {m: i + 1 for i, m in enumerate(
    "enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre".split())}
MESES["setiembre"] = 9
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]


def _fecha_texto(s: str) -> dt.date:
    s = s.strip()
    for fmt in ("%d %b, %Y", "%d %B, %Y", "%d-%b-%y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Fecha no reconocida: {s!r}")


def leer_general(ruta: Path) -> pd.DataFrame:
    ws = openpyxl.load_workbook(ruta, data_only=True).worksheets[0]
    filas = []
    for texto, numero in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        if texto is None or numero is None:
            continue
        parte, dia = str(texto).split(" - ")
        d, _, mes, anio = parte.split()
        f = dt.date(int(anio), MESES[mes.lower()], int(d))
        if DIAS[f.weekday()] != dia.strip().lower():
            raise ValueError(f"Día de semana inconsistente en {texto!r}")
        filas.append((f, int(numero), "Lotería Nacional" if f.weekday() == 6 else "Chances"))
    return pd.DataFrame(filas, columns=["fecha", "numero", "sorteo"]).assign(fuente=ruta.name)


def leer_combinado(ruta: Path, sorteo: str) -> pd.DataFrame:
    ws = openpyxl.load_workbook(ruta, data_only=True).worksheets[0]
    filas = []
    for no, fecha, numero in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        if not isinstance(no, (int, float)) or fecha is None or numero is None:
            continue
        f = fecha.date() if isinstance(fecha, dt.datetime) else _fecha_texto(str(fecha))
        filas.append((f, int(numero), sorteo))
    return pd.DataFrame(filas, columns=["fecha", "numero", "sorteo"]).assign(fuente=ruta.name)


def consolidar(carpeta: Path = Path(".")) -> tuple[pd.DataFrame, pd.DataFrame]:
    gen = leer_general(carpeta / "Histórico general.xlsx")
    lot = leer_combinado(carpeta / "Histórico loteria.xlsx", "Lotería Nacional")
    cha = leer_combinado(carpeta / "Histórico Chances.xlsx", "Chances")
    todo = pd.concat([lot, cha, gen], ignore_index=True)  # prioridad: archivos específicos

    conflictos = (todo.groupby("fecha")
                  .filter(lambda g: g["numero"].nunique() > 1)
                  .sort_values(["fecha", "fuente"]))
    base = todo.drop_duplicates("fecha", keep="first").sort_values("fecha").reset_index(drop=True)
    base["dia"] = pd.to_datetime(base["fecha"]).dt.dayofweek.map(lambda i: DIAS[i])
    return base[["fecha", "dia", "sorteo", "numero", "fuente"]], conflictos
